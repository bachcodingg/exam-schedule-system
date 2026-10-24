"""Read course documents (slides, exercise sheets, past exams) and estimate the workload.

Supported formats: .pdf (needs `pypdf`), .pptx (needs `python-pptx`), .docx (needs
`python-docx`), .txt and .md. Other files (images, archives) are skipped with a note,
and so are files that cannot be read (corrupt, encrypted, or the library is missing).
"""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass, field
from pathlib import Path

from .models import CATEGORIES

# Matched at the start of a word, so "limit" does not match "unlimited".
KEYWORDS = {
    "math_proof": [
        "theorem", "lemma", "proof", "prove", "show that", "corollary", "proposition",
        "q.e.d", "axiom", "if and only if", "subspace", "linearly independent", "basis",
        "eigenv", "kernel", "linear mapping",
    ],
    "math_computational": [
        "integral", "derivative", "differentiat", "integrat", "limit",
        "partial fraction", "chain rule", "taylor", "dx", "∫", "solid of revolution",
        "polar curve", "parametric curve", "arc length",
    ],
    "statistics": [
        "probability", "variance", "distribution", "hypothesis", "regression",
        "confidence interval", "estimator", "random variable", "expected value", "p-value",
        "posterior", "prior", "likelihood", "bayes", "mcmc", "metropolis", "arma",
        "autocorrelation", "time series",
    ],
    "programming": [
        "python", "matlab", "source code", "algorithm", "monte carlo",
        "def ", "import ", "for loop", "compile", "debug",
    ],
    "conceptual": [
        "protocol", "architecture", "agent", "network", "layer", "search algorithm",
        "knowledge base", "routing", "concept",
    ],
    "memorization": ["definition", "terminology", "history", "list of", "acronym"],
}
MINUTES_PER_PAGE = {
    "math_proof": 12, "math_computational": 10, "statistics": 9,
    "programming": 7, "conceptual": 5, "memorization": 6,
}
MINUTES_PER_EXERCISE = {
    "math_proof": 30, "math_computational": 25, "statistics": 20,
    "programming": 30, "conceptual": 12, "memorization": 8,
}
PAST_EXAM_HOURS = 3.0  # time to solve and review one past exam
MAX_PAST_EXAMS = 3
SOLUTION_FACTOR = 0.5  # working through a solved exercise takes about half the time
WORDS_PER_TEXT_PAGE = 400
SUPPORTED = (".pdf", ".pptx", ".docx", ".txt", ".md")

# Task numbers: "Exercise 3", "Problem 2", "Tehtävä 1", or a line starting "4. Let ..." / "4.(a)".
EXERCISE_RE = re.compile(
    r"^\s*(?:(?:exercise|problem|task|question|tehtävä)\s*(\d{1,2})\b|(\d{1,2})\.\s*(?=[(A-ZÄÖ]))",
    re.I | re.M,
)
FORMULA_RE = re.compile(r"[∫∑∏√≤≥≠≈∞∂λμσπ]|\\(frac|int|sum|sqrt)|\^\d|[a-z]_\{?\d")
PAST_EXAM_NAME_RE = re.compile(r"exam|tentti|midterm|final", re.I)
# Exam headers: "Exam 29.4.2025", "Exam, 08.09.2026", "Allowed tools", "Answer all questions".
PAST_EXAM_TEXT_RE = re.compile(
    r"\b(exam|tentti|midterm)\b\W{0,3}\d{1,2}\.\s?\d{1,2}\.\s?\d{4}"
    r"|allowed (printed )?(material|tools)|answer (all|every) questions?|maximum points",
    re.I,
)
REFERENCE_NAME_RE = re.compile(r"sheet|cheat|formula|summary|plan|notes", re.I)
SOLUTION_NAME_RE = re.compile(r"solution|ratkaisu", re.I)


@dataclass
class DocStats:
    path: str
    pages: int
    words: int
    exercises: int
    formula_marks: int
    is_past_exam: bool
    kind: str = "material"  # past exam | solutions | reference | material


@dataclass
class DocEstimate:
    category: str
    hours: float
    scores: dict[str, int]
    docs: list[DocStats] = field(default_factory=list)
    skipped: list[str] = field(default_factory=list)
    unreadable: list[str] = field(default_factory=list)  # "name (reason)"

    @property
    def has_material(self) -> bool:
        """False when only past exams and reference sheets were given (nothing to learn from)."""
        return any(d.kind in ("material", "solutions") for d in self.docs)

    def summary(self) -> str:
        past = sum(d.is_past_exam for d in self.docs)
        tasks = sum(d.exercises for d in self.docs if not d.is_past_exam)
        pages = sum(d.pages for d in self.docs)
        return f"{len(self.docs)} files, {pages} pages, {tasks} exercises, {past} past exams"


def _paginate(text: str) -> list[str]:
    """Cut page-less text into pages of WORDS_PER_TEXT_PAGE words (newlines are kept)."""
    words = text.split(" ")
    return [
        " ".join(words[i : i + WORDS_PER_TEXT_PAGE])
        for i in range(0, max(len(words), 1), WORDS_PER_TEXT_PAGE)
    ]


def extract_pages(path: str | Path) -> list[str]:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        logging.getLogger("pypdf").setLevel(logging.ERROR)  # harmless "wrong pointing object" noise
        return [page.extract_text() or "" for page in PdfReader(path).pages]
    if suffix == ".pptx":
        from pptx import Presentation

        return [
            "\n".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)
            for slide in Presentation(path).slides
        ]
    if suffix in (".txt", ".md"):
        return _paginate(path.read_text(encoding="utf-8", errors="replace"))
    if suffix == ".docx":
        from docx import Document

        doc = Document(path)
        lines = [p.text for p in doc.paragraphs]
        lines += [" | ".join(c.text for c in row.cells) for t in doc.tables for row in t.rows]
        return _paginate("\n".join(lines))
    raise ValueError(f"Unsupported document type: {path.name}")


def classify(text: str) -> tuple[str, dict[str, int]]:
    """Pick the course category whose keywords appear most often."""
    lower = text.lower()
    scores = {
        cat: sum(len(re.findall(r"(?<!\w)" + re.escape(k), lower)) for k in KEYWORDS[cat])
        for cat in CATEGORIES
    }
    best = max(scores, key=lambda c: scores[c])
    return (best if scores[best] > 0 else "conceptual"), scores


def kind_of(path: str | Path, first_page: str) -> str:
    """Past exams are recognised by their file name or by an exam header on the first page."""
    stem = Path(path).stem
    if REFERENCE_NAME_RE.search(stem):
        return "reference"
    if SOLUTION_NAME_RE.search(stem):
        return "solutions"
    if PAST_EXAM_NAME_RE.search(stem) or PAST_EXAM_TEXT_RE.search(first_page):
        return "past exam"
    return "material"


def analyse(path: str | Path) -> tuple[DocStats, str]:
    pages = extract_pages(path)
    text = "\n".join(pages)
    kind = kind_of(path, pages[0] if pages else "")
    # Distinct task numbers, so a bilingual exam or a restated task is not counted twice.
    numbers = {int(a or b) for a, b in EXERCISE_RE.findall(text)}
    stats = DocStats(
        path=str(path),
        pages=len(pages),
        words=len(text.split()),
        exercises=0 if kind == "reference" else len(numbers),
        formula_marks=len(FORMULA_RE.findall(text)),
        is_past_exam=kind == "past exam",
        kind=kind,
    )
    return stats, text


def estimate_documents(paths: list[str | Path], category: str | None = None) -> DocEstimate:
    """Estimate study hours from the material. `category` overrides the automatic guess."""
    readable = [p for p in paths if Path(p).suffix.lower() in SUPPORTED]
    skipped = [Path(p).name for p in paths if Path(p).suffix.lower() not in SUPPORTED]
    analysed, unreadable = [], []
    for p in readable:
        try:
            analysed.append(analyse(p))
        except ImportError as err:
            unreadable.append(f"{Path(p).name} (install {err.name or 'the reader library'})")
        except Exception as err:  # corrupt or encrypted files should not stop the whole plan
            unreadable.append(f"{Path(p).name} ({type(err).__name__})")
    all_text = "\n".join(text for _, text in analysed)
    guessed, scores = classify(all_text)
    cat = category or guessed

    minutes = 0.0
    past_exams = 0
    for stats, _ in analysed:
        if stats.is_past_exam:
            past_exams += 1
            continue
        exercise_minutes = stats.exercises * MINUTES_PER_EXERCISE[cat]
        if stats.kind == "solutions" and stats.exercises:
            # The task sheet is counted on its own; here you only work through the solutions.
            minutes += exercise_minutes * SOLUTION_FACTOR
            continue
        page_minutes = stats.pages * MINUTES_PER_PAGE[cat]
        if stats.pages and stats.formula_marks / stats.pages > 5:
            page_minutes *= 1.15  # formula-dense material reads slower
        minutes += page_minutes + exercise_minutes
    hours = minutes / 60 + min(past_exams, MAX_PAST_EXAMS) * PAST_EXAM_HOURS

    return DocEstimate(cat, round(hours, 1), scores, [s for s, _ in analysed], skipped, unreadable)
